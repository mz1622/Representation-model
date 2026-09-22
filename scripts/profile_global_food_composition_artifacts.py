#!/usr/bin/env python3
"""Profile raw artifacts collected for the global FCDB landscape review.

The output is descriptive only. It does not normalize units, map food or
component identities, remove duplicates, or produce a model-ready table.
"""

from __future__ import annotations

import csv
import hashlib
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
INVENTORY = ROOT / "data" / "raw" / "global_fcdb_inventory_2026_09_10"
DOWNLOADS = INVENTORY / "downloads"
OUT = INVENTORY / "artifact_profile.csv"


def digest(path: Path) -> str:
    hasher = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            hasher.update(block)
    return hasher.hexdigest()


def count_ids(path: Path, sheet_names: list[str], pattern: str) -> int:
    matcher = re.compile(pattern)
    values: set[str] = set()
    workbook = load_workbook(path, read_only=True, data_only=True)
    for sheet_name in sheet_names:
        sheet = workbook[sheet_name]
        for (value,) in sheet.iter_rows(min_col=1, max_col=1, values_only=True):
            if isinstance(value, str) and matcher.fullmatch(value.strip()):
                values.add(value.strip())
    return len(values)


def workbook_shape(path: Path, sheet_name: str) -> tuple[int, int]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[sheet_name]
    return sheet.max_row, sheet.max_column


def profile(path: Path, source_key: str, structure: str, food_count: str, component_count: str, readiness: str) -> dict[str, str]:
    return {
        "source_key": source_key,
        "artifact": str(path.relative_to(ROOT)),
        "format": path.suffix.lower().lstrip("."),
        "bytes": str(path.stat().st_size),
        "sha256": digest(path),
        "observed_food_profile_count": food_count,
        "observed_component_count": component_count,
        "raw_structure": structure,
        "ingestion_readiness": readiness,
    }


def main() -> None:
    records: list[dict[str, str]] = []

    wafct = DOWNLOADS / "wafct_2019.xlsx"
    wafct_count = count_ids(wafct, ["03 NV_sum_39 (per 100g EP)"], r"\d{2}_\d{3}")
    records.append(profile(
        wafct,
        "wafct_2019",
        "Summary sheets have 39- and 57-component panels plus bibliography, recipe, yield and retention sheets. "
        "The 39-component summary marks 99 of 1,028 food codes as 'calc. from recipe'.",
        str(wafct_count),
        "57 main component definitions in the Components sheet",
        "Machine-readable; future construction must split recipe-derived dishes from source-linked profiles.",
    ))

    lesotho = DOWNLOADS / "lesotho_fct_2006.xlsm"
    lesotho_count = count_ids(lesotho, ["source DB original"], r"\d{6}")
    records.append(profile(
        lesotho,
        "lesotho_fct_2006",
        "Macro-enabled workbook with original-source, reference, recipe and printing sheets.",
        str(lesotho_count),
        "119 source-DB columns (metadata and composition together)",
        "Machine-readable, but historical and recipe/reference structure requires a value-level audit.",
    ))

    anfood = DOWNLOADS / "anfood_2_0.xlsx"
    anfood_sheets = [name for name in load_workbook(anfood, read_only=True).sheetnames if name[:2].isdigit()]
    anfood_count = count_ids(anfood, anfood_sheets, r"\d{7}")
    records.append(profile(
        anfood,
        "anfood_2_0",
        "Food-group sheets retain food item ID, country/region, processing, scientific name, sample count, method comments and bibliography.",
        str(anfood_count),
        "up to 245 columns on food-group sheets",
        "High-value analytical layer; do not collapse literature observations with different sample facets.",
    ))

    biofood = DOWNLOADS / "biofoodcomp_4_0.xlsx"
    biofood_sheets = [name for name in load_workbook(biofood, read_only=True).sheetnames if name[:2].isdigit()]
    biofood_count = count_ids(biofood, biofood_sheets, r"\d{7}")
    records.append(profile(
        biofood,
        "biofoodcomp_4_0",
        "Food-group sheets retain country, cultivar/variety/breed, processing, sample count, method comments and bibliography.",
        str(biofood_count),
        "up to 247 columns on food-group sheets",
        "High-value biodiversity layer; varieties and locations are intentional distinct observations, not duplicate names.",
    ))

    phyfood = DOWNLOADS / "phyfoodcomp_1_0.xlsx"
    phyfood_sheets = [name for name in load_workbook(phyfood, read_only=True).sheetnames if name[:2].isdigit()]
    phyfood_count = count_ids(phyfood, phyfood_sheets, r"\d{8}")
    records.append(profile(
        phyfood,
        "phyfoodcomp_1_0",
        "Food-group sheets provide FoodEx2 descriptions and phytate/mineral-related fields alongside provenance and processing facets.",
        str(phyfood_count),
        "35 component definitions in the Components sheet",
        "Specialist phytate/bioavailability layer; not a general nutrient corpus.",
    ))

    swiss = DOWNLOADS / "swiss_fcdb_v7_1.xlsx"
    generic_rows, generic_cols = workbook_shape(swiss, "Aliments génériques")
    branded_rows, branded_cols = workbook_shape(swiss, "Produits de marque")
    records.append(profile(
        swiss,
        "swiss_fcdb_7_1",
        f"Generic table is {generic_rows} x {generic_cols}; branded-product table is {branded_rows} x {branded_cols}; each has three heading rows.",
        f"{generic_rows - 3} generic rows; {branded_rows - 3} branded-product rows",
        "Nutrient-code sheet included; food sheets have 144 columns",
        "Machine-readable; generic and branded foods should remain separate future layers.",
    ))

    bangladesh = DOWNLOADS / "bangladesh_fct_2013.xlsx"
    bangladesh_count = count_ids(bangladesh, ["UserDB_Main_table"], r"\d{2}_\d{4}")
    records.append(profile(
        bangladesh,
        "bangladesh_fct_2013",
        "Main user table has component, bibliography, yield/retention and separate amino-acid, fatty-acid, antioxidant and antinutrient annexes. "
        "Rows labelled Recipe calculation and source IDs remain visible.",
        str(bangladesh_count),
        "244 component rows in the Components sheet",
        "Machine-readable; later use must preserve recipe/source flags and annex-specific definitions.",
    ))

    smiling = [
        ("smiling_cambodia_2013", "smiling_cambodia_values.xlsx", "90 official user-table foods; paired value-reference and quality-assessment workbooks."),
        ("smiling_indonesia_2013", "smiling_indonesia_values.xlsx", "174 official user-table foods; paired value-reference and quality-assessment workbooks."),
        ("smiling_laos_2013", "smiling_laos_values.xlsx", "140 official user-table foods; paired value-reference and quality-assessment workbooks."),
        ("smiling_thailand_2013", "smiling_thailand_values.xlsx", "142 official user-table foods; paired value-reference and quality-assessment workbooks."),
        ("smiling_vietnam_2013", "smiling_vietnam_values.xlsx", "162 official user-table foods; paired value-reference and quality-assessment workbooks."),
    ]
    for source_key, filename, structure in smiling:
        path = DOWNLOADS / filename
        records.append(profile(
            path,
            source_key,
            structure,
            structure.split(" official")[0],
            "source-specific component definitions and quality workbook",
            "Machine-readable and value-level documented; small, historical regional supplements pending original-source lineage review.",
        ))

    for source_key, filename, readiness in [
        ("kenya_fct_2018", "kenya_fct_2018.pdf", "PDF only; seek native table or conduct audited double extraction before any data use."),
        ("malawi_fct_2019", "malawi_fct_2019.pdf", "PDF only; seek native table or conduct audited double extraction before any data use."),
        ("ethiopian_fct_2025", "ethiopian_fct_2025_user_guide.pdf", "Condensed official user-guide PDF only; acquire the 69-component value-documented Excel datasheets before any numeric ingestion."),
    ]:
        path = DOWNLOADS / filename
        records.append(profile(
            path,
            source_key,
            "Official table retained as PDF; no machine-readable structure was inferred.",
            "not extracted", "not extracted", readiness,
        ))

    fields = list(records[0])
    with OUT.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(records)
    print(f"Wrote {len(records)} raw-artifact profiles to {OUT}")


if __name__ == "__main__":
    main()
