"""Reproducible scoping-review and source-registry builder.

The automated search creates an auditable candidate corpus. It deliberately
does not pretend to replace dual-reviewer title/full-text screening.
"""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import re
import time
import urllib.parse
import urllib.request
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .constants import AS_OF_DATE, DATASET_VERSION, FAO_SCREENING_QUESTIONS
from .util import normalize_text, sha256_file, stable_id, write_csv, write_json


SEARCH_QUERIES = [
    ("food composition database review", '"food composition database" AND (review OR landscape OR inventory)'),
    ("database quality", '"food composition database" AND (quality OR evaluation OR EuroFIR OR INFOODS)'),
    ("food matching", '(food matching OR food linkage OR food entity linking) AND (composition OR nutrient)'),
    ("component harmonization", '(nutrient OR food component) AND (harmonization OR standardization) AND database'),
    ("unit conversion", '(food composition OR nutrient database) AND (unit conversion OR expression basis OR dry weight)'),
    ("food ontology", '(FoodOn OR LanguaL OR FoodEx2) AND (food composition OR nutrition)'),
    ("component ontology", '(INFOODS OR CDNO OR ChEBI OR "LIPID MAPS") AND food composition'),
    ("prediction benchmark", '(food composition OR nutrient) AND (prediction OR imputation) AND benchmark'),
]


CURATED_REFERENCES = [
    {
        "record_id": "doi:10.7326/M18-0850",
        "title": "PRISMA extension for scoping reviews (PRISMA-ScR): checklist and explanation",
        "year": 2018,
        "url": "https://doi.org/10.7326/M18-0850",
        "record_type": "peer_reviewed_reporting_guideline",
        "decision": "include",
        "reason": "Reporting framework for the scoping review.",
    },
    {
        "record_id": "doi:10.1186/s13643-020-01542-z",
        "title": "PRISMA-S: an extension to the PRISMA Statement for Reporting Literature Searches in Systematic Reviews",
        "year": 2021,
        "url": "https://doi.org/10.1186/s13643-020-01542-z",
        "record_type": "peer_reviewed_reporting_guideline",
        "decision": "include",
        "reason": "Reporting framework for reproducible database, website and citation searches.",
    },
    {
        "record_id": "doi:10.3389/fnut.2025.1552367",
        "title": "Global landscape assessment of food composition databases",
        "year": 2025,
        "url": "https://doi.org/10.3389/fnut.2025.1552367",
        "record_type": "peer_reviewed_review",
        "decision": "include",
        "reason": "Global review and FAIR assessment of 101 included FCDBs; seeds the source registry.",
    },
    {
        "record_id": "fao:cc5371en",
        "title": "FAO/INFOODS Evaluation framework to assess the quality of published food composition tables and databases",
        "year": 2023,
        "url": "https://doi.org/10.4060/cc5371en",
        "record_type": "official_standard",
        "decision": "include",
        "reason": "Primary source-screening and database-evaluation framework.",
    },
    {
        "record_id": "doi:10.1016/j.foodchem.2015.02.110",
        "title": "EuroFIR quality approach for managing food composition data; where are we in 2014?",
        "year": 2016,
        "url": "https://doi.org/10.1016/j.foodchem.2015.02.110",
        "record_type": "peer_reviewed_method",
        "decision": "include",
        "reason": "Describes the EuroFIR quality-management framework and value documentation process.",
    },
    {
        "record_id": "eurofir:qe-scirep-2009",
        "title": "Guidelines for quality index attribution to original data from scientific literature or reports for EuroFIR data interchange",
        "year": 2009,
        "url": "https://www.fao.org/infoods/infoods/training/documents-used-in-food-composition-study-guide/en/",
        "record_type": "technical_quality_guideline",
        "decision": "include",
        "reason": "Primary QE-SCIREP criteria for food description, component identity, sampling, sample handling, analytical method and analytical quality control.",
    },
    {
        "record_id": "fao:food-matching-v1.2",
        "title": "FAO/INFOODS Guidelines for Food Matching, Version 1.2",
        "year": 2012,
        "url": "https://www.fao.org/4/ap805e/ap805e.pdf",
        "record_type": "official_standard",
        "decision": "include",
        "reason": "Defines the food characteristics and documentation needed for defensible matching.",
    },
    {
        "record_id": "fao:unit-conversion-v1.0",
        "title": "FAO/INFOODS Guidelines for Converting Units, Denominators and Expressions, Version 1.0",
        "year": 2012,
        "url": "https://www.fao.org/infoods/infoods/standards-guidelines/en/",
        "record_type": "official_standard",
        "decision": "include",
        "reason": "Controls unit, denominator and expression conversion into the main mass modality.",
    },
    {
        "record_id": "fao:data-checking-v1.0",
        "title": "FAO/INFOODS Guidelines for Checking Food Composition Data prior to Publication of a User Table/Database, Version 1.0",
        "year": 2012,
        "url": "https://www.fao.org/infoods/infoods/standards-guidelines/en/",
        "record_type": "official_standard",
        "decision": "include",
        "reason": "Basis for automated and expert composition plausibility checks.",
    },
    {
        "record_id": "doi:10.1038/sdata.2016.18",
        "title": "The FAIR Guiding Principles for scientific data management and stewardship",
        "year": 2016,
        "url": "https://doi.org/10.1038/sdata.2016.18",
        "record_type": "peer_reviewed_standard",
        "decision": "include",
        "reason": "Defines provenance, persistent identification and reusable-release requirements.",
    },
    {
        "record_id": "doi:10.1002/csc2.20092",
        "title": "Knowledge representation and data sharing to unlock crop variation for nutritional food security",
        "year": 2020,
        "url": "https://doi.org/10.1002/csc2.20092",
        "record_type": "peer_reviewed_ontology_method",
        "decision": "include",
        "reason": "Scientific basis for CDNO component terminology and hierarchy.",
    },
    {
        "record_id": "foodon:official",
        "title": "FoodOn food ontology",
        "year": 2026,
        "url": "https://foodon.org/",
        "record_type": "official_ontology",
        "decision": "include",
        "reason": "Food identity, process and material hierarchy authority.",
    },
    {
        "record_id": "chebi:official",
        "title": "Chemical Entities of Biological Interest (ChEBI)",
        "year": 2026,
        "url": "https://www.ebi.ac.uk/chebi/",
        "record_type": "official_ontology",
        "decision": "include",
        "reason": "Chemical identity, preferred names and chemical-class hierarchy.",
    },
    {
        "record_id": "lipidmaps:classification",
        "title": "LIPID MAPS Lipid Classification System",
        "year": 2024,
        "url": "https://www.lipidmaps.org/resources/education/classification",
        "record_type": "official_classification",
        "decision": "include",
        "reason": "Authoritative lipid category, class and subclass system.",
    },
]


CITATION_TRACKING_SEEDS = {
    "doi:10.3389/fnut.2025.1552367",
    "doi:10.1016/j.foodchem.2015.02.110",
    "doi:10.1002/csc2.20092",
}


SOURCE_OVERRIDES = {
    "USDA National Nutrient Database for Standard Reference Legacy Release": {
        "source_key": "usda_sr_legacy", "local_status": "available", "license": "United States public domain",
        "version": "April 2018 / FDC export 2019", "layer": "primary_reference",
        "official_url": "https://fdc.nal.usda.gov/data-documentation.html",
    },
    "Canadian Nutrient File": {
        "source_key": "cnf", "local_status": "available", "license": "Open Government Licence - Canada",
        "version": "2026", "layer": "primary_reference",
        "official_url": "https://food-nutrition.canada.ca/cnf-fce/",
    },
    "Fødevaredata - Frida Food Data": {
        "source_key": "frida", "local_status": "available", "license": "CC BY 4.0",
        "version": "6.1 (2026)", "layer": "primary_reference",
        "official_url": "https://doi.org/10.11583/DTU.32312844.v1",
    },
    "Ciqual - French Food Composition Table": {
        "source_key": "ciqual", "local_status": "available", "license": "CC BY 4.0",
        "version": "2025", "layer": "primary_reference",
        "official_url": "https://doi.org/10.5281/zenodo.17550133",
    },
    "Australian Food Composition Database": {
        "source_key": "afcd", "local_status": "available", "license": "CC BY-SA 3.0 AU",
        "version": "Release 3", "layer": "primary_reference",
        "official_url": "https://www.foodstandards.gov.au/science-data/monitoringnutrients/afcd",
    },
    "Composition of Foods Integrated Dataset": {
        "source_key": "cofid", "local_status": "available", "license": "UK Open Government Licence",
        "version": "2021", "layer": "primary_reference",
        "official_url": "https://www.gov.uk/government/publications/composition-of-foods-integrated-dataset-cofid",
    },
    "Finnish National Food Composition Database": {
        "source_key": "fineli", "local_status": "download_blocked", "license": "CC BY 4.0",
        "version": "current site as of 2026-09-03", "layer": "candidate_primary",
        "official_url": "https://fineli.fi/fineli/en/avoin-data",
    },
    "New Zealand Food Composition Database": {
        "source_key": "foodfiles", "local_status": "not_ingested", "license": "terms require review",
        "version": "FOODfiles 2024", "layer": "candidate_primary",
        "official_url": "https://www.foodcomposition.co.nz/foodfiles/",
    },
}


LOCAL_SOURCE_FILES = {
    "usda_sr_legacy": "data/raw/usda/sr_legacy_2018",
    "usda_foundation": "data/raw/usda/foundation_2026_04_30",
    "fndds": "data/raw/usda/fndds_2021_2023",
    "cnf": "data/raw/cnf_2026",
    "frida": "data/raw/candidates/frida_6_1",
    "ciqual": "data/raw/candidates/ciqual_2025",
    "afcd": "data/raw/candidates/afcd_release_3",
    "cofid": "data/raw/candidates/cofid_2021",
    "foodb": "foodb_2020_04_07_csv",
}


PRELIMINARY_SCREENING = {
    "usda_sr_legacy": ("yes", "yes", "no", "yes", "yes", "yes", "yes", "yes"),
    "usda_foundation": ("yes", "yes", "no", "yes", "yes", "yes", "yes", "yes"),
    "fndds": ("yes", "yes", "no", "yes", "yes", "yes", "yes", "yes"),
    "cnf": ("yes", "yes", "yes", "yes", "yes", "yes", "yes", "yes"),
    "frida": ("yes", "yes", "yes", "yes", "yes", "yes", "yes", "yes"),
    "ciqual": ("yes", "yes", "yes", "yes", "yes", "yes", "yes", "yes"),
    "afcd": ("yes", "yes", "yes", "yes", "yes", "yes", "yes", "yes"),
    "cofid": ("yes", "yes", "yes", "yes", "yes", "yes", "yes", "yes"),
    "foodb": ("no", "no", "no", "yes", "yes", "yes", "yes", "yes"),
}


def _request_json(url: str, cache_path: Path, offline: bool) -> dict[str, Any]:
    if cache_path.exists():
        return json.loads(cache_path.read_text())
    if offline:
        raise FileNotFoundError(f"Offline cache is missing: {cache_path}")
    request = urllib.request.Request(url, headers={"User-Agent": "FoodNutritionScopingReview/1.0 (research use)"})
    with urllib.request.urlopen(request, timeout=60) as response:
        payload = json.load(response)
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    cache_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2))
    time.sleep(0.34)
    return payload


def build_fao_global_directory(out_dir: Path, offline: bool) -> pd.DataFrame:
    """Snapshot every entry exposed by the official FAO global directory."""
    cache = out_dir / "api_cache/fao_global_directory_2026-09-03.html"
    if cache.exists():
        page = cache.read_text(encoding="utf-8", errors="replace")
    elif offline:
        raise FileNotFoundError(f"Offline FAO directory cache is missing: {cache}")
    else:
        request = urllib.request.Request(
            "https://www.fao.org/food-composition/tables-and-databases/1/en",
            headers={"User-Agent": "FoodNutritionScopingReview/1.0 (research use)"},
        )
        with urllib.request.urlopen(request, timeout=60) as response:
            page = response.read().decode("utf-8", errors="replace")
        cache.parent.mkdir(parents=True, exist_ok=True)
        cache.write_text(page, encoding="utf-8")
    match = re.search(r"const\s+items\s*=\s*(\[.*?\]);", page, flags=re.DOTALL)
    if not match:
        raise ValueError("Could not locate the official FAO directory item payload.")
    payload = json.loads(html.unescape(match.group(1)))
    rows = []
    for item in payload:
        locations = item.get("locations") or []
        rows.append({
            "fao_directory_id": stable_id("fao-directory", item.get("url", ""), item.get("title", "")),
            "title": item.get("title", ""),
            "official_url": item.get("url", ""),
            "countries": ";".join(sorted({str(location.get("country", "")) for location in locations if location.get("country")})),
            "iso3": ";".join(sorted({str(location.get("iso3", "")) for location in locations if location.get("iso3")})),
            "directory_snapshot_date": AS_OF_DATE,
            "inclusion_is_fao_endorsement": False,
        })
    directory = pd.DataFrame(rows).drop_duplicates("fao_directory_id")
    if directory.empty:
        raise ValueError("The official FAO directory snapshot contained no entries.")
    return directory


def _search_crossref(query: str, cache: Path, offline: bool) -> tuple[int, list[dict[str, Any]]]:
    url = "https://api.crossref.org/works?" + urllib.parse.urlencode({
        "query.bibliographic": query,
        "filter": "from-pub-date:2000-01-01,until-pub-date:2026-09-03",
        "rows": 100,
        "select": "DOI,title,published,URL,type,publisher",
    })
    data = _request_json(url, cache, offline)["message"]
    rows = []
    for item in data.get("items", []):
        title = (item.get("title") or [""])[0]
        year_parts = (item.get("published") or {}).get("date-parts") or [[None]]
        doi = (item.get("DOI") or "").lower()
        rows.append({
            "record_id": f"doi:{doi}" if doi else stable_id("crossref", title),
            "title": title, "year": year_parts[0][0], "doi": doi,
            "url": item.get("URL", ""), "provider": "Crossref",
            "record_type": item.get("type", ""), "publisher": item.get("publisher", ""),
        })
    return int(data.get("total-results", len(rows))), rows


def _search_openalex(query: str, cache: Path, offline: bool) -> tuple[int, list[dict[str, Any]]]:
    url = "https://api.openalex.org/works?" + urllib.parse.urlencode({
        "search": query, "per-page": 100,
        "filter": "from_publication_date:2000-01-01,to_publication_date:2026-09-03",
    })
    data = _request_json(url, cache, offline)
    rows = []
    for item in data.get("results", []):
        doi_url = item.get("doi") or ""
        doi = doi_url.removeprefix("https://doi.org/").lower()
        rows.append({
            "record_id": f"doi:{doi}" if doi else item.get("id", stable_id("openalex", item.get("title"))),
            "title": item.get("title", ""), "year": item.get("publication_year"), "doi": doi,
            "url": (item.get("primary_location") or {}).get("landing_page_url") or item.get("id", ""),
            "provider": "OpenAlex", "record_type": item.get("type", ""),
            "publisher": (((item.get("primary_location") or {}).get("source") or {}).get("display_name") or ""),
        })
    return int((data.get("meta") or {}).get("count", len(rows))), rows


def _search_pubmed(query: str, cache: Path, offline: bool) -> tuple[int, list[dict[str, Any]]]:
    encoded = urllib.parse.urlencode({"db": "pubmed", "term": query, "retmode": "json", "retmax": 100, "mindate": "2000", "maxdate": "2026/09/03"})
    search = _request_json(f"https://eutils.ncbi.nlm.nih.gov/entrez/eutils/esearch.fcgi?{encoded}", cache.with_name(cache.stem + "_search.json"), offline)
    ids = search.get("esearchresult", {}).get("idlist", [])
    if not ids:
        return int(search.get("esearchresult", {}).get("count", 0)), []
    summary_url = "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/esummary.fcgi?" + urllib.parse.urlencode({"db": "pubmed", "id": ",".join(ids), "retmode": "json"})
    data = _request_json(summary_url, cache.with_name(cache.stem + "_summary.json"), offline).get("result", {})
    rows = []
    for pmid in ids:
        item = data.get(pmid, {})
        article_ids = {x.get("idtype"): x.get("value") for x in item.get("articleids", [])}
        doi = (article_ids.get("doi") or "").lower()
        year_text = (item.get("pubdate") or "")[:4]
        rows.append({
            "record_id": f"doi:{doi}" if doi else f"pmid:{pmid}", "title": item.get("title", ""),
            "year": int(year_text) if year_text.isdigit() else None, "doi": doi,
            "url": f"https://pubmed.ncbi.nlm.nih.gov/{pmid}/", "provider": "PubMed",
            "record_type": "journal_article", "publisher": item.get("fulljournalname", ""), "pmid": pmid,
        })
    return int(search.get("esearchresult", {}).get("count", len(rows))), rows


def build_search_corpus(out_dir: Path, offline: bool) -> tuple[pd.DataFrame, pd.DataFrame]:
    cache_dir = out_dir / "api_cache"
    logs: list[dict[str, Any]] = []
    records: list[dict[str, Any]] = []
    searchers = (("Crossref", _search_crossref), ("OpenAlex", _search_openalex), ("PubMed", _search_pubmed))
    for query_index, (topic, query) in enumerate(SEARCH_QUERIES, start=1):
        for provider, searcher in searchers:
            cache = cache_dir / f"q{query_index:02d}_{provider.casefold()}.json"
            try:
                total, found = searcher(query, cache, offline)
                status, error = "completed", ""
            except (OSError, ValueError, KeyError) as exc:
                total, found, status, error = 0, [], "failed", str(exc)
            search_id = f"Q{query_index:02d}-{provider}"
            logs.append({"search_id": search_id, "topic": topic, "provider": provider, "query": query, "search_date": AS_OF_DATE, "total_results": total, "retrieved": len(found), "status": status, "error": error})
            for row in found:
                records.append({**row, "search_id": search_id, "topic": topic})
    record_df = pd.DataFrame(records)
    if not record_df.empty:
        record_df["normalized_title"] = record_df["title"].map(normalize_text)
        record_df["dedup_key"] = record_df.apply(lambda r: r["record_id"] if str(r["record_id"]).startswith(("doi:", "pmid:")) else stable_id("title", r["normalized_title"]), axis=1)
        record_df["screening_status"] = "awaiting_dual_reviewer_screening"
        record_df["screening_reason"] = "Automated search result; no inclusion claim until human screening."
    return pd.DataFrame(logs), record_df


def _openalex_row(item: dict[str, Any], *, seed_record_id: str, direction: str) -> dict[str, Any]:
    doi_url = item.get("doi") or ""
    doi = doi_url.removeprefix("https://doi.org/").lower()
    title = item.get("title", "")
    return {
        "record_id": f"doi:{doi}" if doi else item.get("id", stable_id("openalex", title)),
        "title": title,
        "year": item.get("publication_year"),
        "doi": doi,
        "url": (item.get("primary_location") or {}).get("landing_page_url") or item.get("id", ""),
        "provider": "OpenAlex citation index",
        "record_type": item.get("type", ""),
        "publisher": (((item.get("primary_location") or {}).get("source") or {}).get("display_name") or ""),
        "search_id": f"citation:{seed_record_id}:{direction}",
        "topic": "backward and forward citation tracking",
        "citation_seed_record_id": seed_record_id,
        "citation_direction": direction,
        "openalex_id": item.get("id", ""),
    }


def build_citation_tracking(out_dir: Path, offline: bool) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Retrieve complete OpenAlex backward/forward sets for domain-method seeds."""
    cache_dir = out_dir / "api_cache/citation_tracking"
    logs: list[dict[str, Any]] = []
    records: list[dict[str, Any]] = []
    for seed_record_id in sorted(CITATION_TRACKING_SEEDS):
        doi = seed_record_id.removeprefix("doi:")
        seed_url = "https://api.openalex.org/works?" + urllib.parse.urlencode({
            "filter": f"doi:{doi}", "per-page": 1,
        })
        seed_payload = _request_json(seed_url, cache_dir / f"{hashlib.sha256(seed_record_id.encode()).hexdigest()[:12]}_seed.json", offline)
        seed_results = seed_payload.get("results", [])
        if not seed_results:
            logs.append({
                "search_id": f"citation:{seed_record_id}:seed", "topic": "citation tracking",
                "provider": "OpenAlex", "query": f"doi:{doi}", "search_date": AS_OF_DATE,
                "total_results": 0, "retrieved": 0, "status": "failed", "error": "Seed DOI not found",
            })
            continue
        seed = seed_results[0]
        openalex_id = str(seed.get("id", "")).rsplit("/", 1)[-1]

        backward_ids = [str(value).rsplit("/", 1)[-1] for value in seed.get("referenced_works", []) if value]
        backward_rows: list[dict[str, Any]] = []
        for chunk_index in range(0, len(backward_ids), 50):
            chunk = backward_ids[chunk_index:chunk_index + 50]
            query_url = "https://api.openalex.org/works?" + urllib.parse.urlencode({
                "filter": "openalex_id:" + "|".join(chunk), "per-page": 200,
            })
            payload = _request_json(
                query_url,
                cache_dir / f"{openalex_id}_backward_{chunk_index // 50 + 1:03d}.json",
                offline,
            )
            backward_rows.extend(payload.get("results", []))
        records.extend(_openalex_row(item, seed_record_id=seed_record_id, direction="backward") for item in backward_rows)
        logs.append({
            "search_id": f"citation:{seed_record_id}:backward", "topic": "backward citation tracking",
            "provider": "OpenAlex", "query": f"referenced_works of {openalex_id}", "search_date": AS_OF_DATE,
            "total_results": len(backward_ids), "retrieved": len(backward_rows),
            "status": "completed" if len(backward_rows) == len(backward_ids) else "incomplete",
            "error": "" if len(backward_rows) == len(backward_ids) else "OpenAlex did not return every referenced-work ID",
        })

        forward_rows: list[dict[str, Any]] = []
        cursor = "*"
        page = 1
        expected = None
        while cursor:
            query_url = "https://api.openalex.org/works?" + urllib.parse.urlencode({
                "filter": f"cites:{openalex_id},to_publication_date:{AS_OF_DATE}",
                "per-page": 200, "cursor": cursor,
            })
            payload = _request_json(query_url, cache_dir / f"{openalex_id}_forward_{page:03d}.json", offline)
            if expected is None:
                expected = int((payload.get("meta") or {}).get("count", 0))
            page_rows = payload.get("results", [])
            forward_rows.extend(page_rows)
            cursor = (payload.get("meta") or {}).get("next_cursor") if page_rows else None
            page += 1
        records.extend(_openalex_row(item, seed_record_id=seed_record_id, direction="forward") for item in forward_rows)
        logs.append({
            "search_id": f"citation:{seed_record_id}:forward", "topic": "forward citation tracking",
            "provider": "OpenAlex", "query": f"cites:{openalex_id}; through {AS_OF_DATE}", "search_date": AS_OF_DATE,
            "total_results": expected or 0, "retrieved": len(forward_rows),
            "status": "completed" if len(forward_rows) == (expected or 0) else "incomplete",
            "error": "" if len(forward_rows) == (expected or 0) else "Retrieved count differs from OpenAlex count",
        })

    record_df = pd.DataFrame(records)
    if not record_df.empty:
        record_df["normalized_title"] = record_df["title"].map(normalize_text)
        record_df["dedup_key"] = record_df.apply(
            lambda row: row["record_id"] if str(row["record_id"]).startswith(("doi:", "pmid:")) else stable_id("title", row["normalized_title"]),
            axis=1,
        )
        record_df["screening_status"] = "awaiting_dual_reviewer_screening"
        record_df["screening_reason"] = "Citation-index candidate; no inclusion claim until human screening."
    return pd.DataFrame(logs), record_df


def build_source_registry(root: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    workbook = root / "data/raw/reference/frontiers_2025_fcdb_review/PTFI-FS-FCDBs-LandscapeAssessment-Supplementary-2.0.0_20250220.xlsx"
    general = pd.read_excel(workbook, sheet_name="Supp. Table S1. General Data")
    nutrition = pd.read_excel(workbook, sheet_name="Supp. Table S2. Food and Nutrit")
    registry = general.merge(nutrition, on="Database Name", how="left")
    registry.columns = [normalize_text(c).replace(" ", "_") for c in registry.columns]
    registry.insert(0, "registry_id", [stable_id("fcdb", x) for x in registry["database_name"]])
    registry["review_origin"] = "Brinkley et al. 2025 supplementary dataset"
    registry["source_key"] = ""
    registry["local_status"] = "registry_only"
    registry["layer"] = "candidate_unreviewed"
    registry["license"] = registry.get("licensing", "")
    registry["version"] = ""
    registry["official_url"] = registry.get("link", "")
    for database_name, override in SOURCE_OVERRIDES.items():
        mask = registry["database_name"].map(normalize_text).eq(normalize_text(database_name))
        if mask.any():
            for key, value in override.items():
                registry.loc[mask, key] = value

    explicit = [
        ("USDA National Nutrient Database for Standard Reference Legacy Release", "usda_sr_legacy", "primary_reference", "available", "United States public domain", "April 2018 / FDC export 2019", "https://fdc.nal.usda.gov/data-documentation.html"),
        ("USDA Foundation Foods", "usda_foundation", "locked_source_validation", "available", "United States public domain", "April 2026", "https://fdc.nal.usda.gov/data-documentation.html"),
        ("USDA Food and Nutrient Database for Dietary Studies", "fndds", "calculated_dish_auxiliary", "available", "United States public domain", "2021-2023", "https://fdc.nal.usda.gov/fndds.html"),
        ("FooDB", "foodb", "specialist_and_lineage_archive", "available", "FooDB terms require redistribution review", "1.0 (2020 export)", "https://foodb.ca/"),
    ]
    for name, key, layer, status, license_text, version, url in explicit:
        if not registry["source_key"].eq(key).any():
            blank = {column: "" for column in registry.columns}
            blank.update({"registry_id": stable_id("fcdb", name), "database_name": name, "source_key": key, "layer": layer, "local_status": status, "license": license_text, "version": version, "official_url": url, "review_origin": "project source registry"})
            registry = pd.concat([registry, pd.DataFrame([blank])], ignore_index=True)

    registry["build_decision"] = np.select(
        [registry["local_status"].eq("available") & registry["layer"].isin(["primary_reference", "locked_source_validation"]),
         registry["local_status"].eq("available") & registry["layer"].eq("calculated_dish_auxiliary"),
         registry["local_status"].eq("available") & registry["layer"].eq("specialist_and_lineage_archive")],
        ["candidate_main_or_locked_validation", "auxiliary_only", "specialist_or_provenance_only"],
        default="not_ingested_in_v1",
    )
    registry["build_decision_reason"] = np.select(
        [registry["build_decision"].eq("candidate_main_or_locked_validation"),
         registry["build_decision"].eq("auxiliary_only"),
         registry["build_decision"].eq("specialist_or_provenance_only"),
         registry["local_status"].eq("download_blocked")],
        ["Structured version and licence are locally available; individual values still pass value-level quality gates.",
         "Survey or calculated dishes are retained for auxiliary training but excluded from primary validation.",
         "Values are decomposed by original citation lineage and are not counted as independent general-purpose evidence.",
         "Official source is identified, but reproducible automated download was blocked at build time."],
        default="Registry entry retained for review; version, licence, acquisition or full evaluation is incomplete.",
    )

    files = []
    for source_key, relative in LOCAL_SOURCE_FILES.items():
        base = root / relative
        if not base.exists():
            continue
        for path in sorted(p for p in base.rglob("*") if p.is_file() and not p.name.startswith(".DS_Store")):
            files.append({"source_key": source_key, "relative_path": str(path.relative_to(root)), "bytes": path.stat().st_size, "sha256": sha256_file(path)})
    for source_key, relative in (("authority_reference", "data/raw/reference"), ("ontology_snapshot", "data/raw/ontology")):
        base = root / relative
        if not base.exists():
            continue
        for path in sorted(p for p in base.rglob("*") if p.is_file() and not p.name.startswith(".DS_Store")):
            files.append({"source_key": source_key, "relative_path": str(path.relative_to(root)), "bytes": path.stat().st_size, "sha256": sha256_file(path)})
    return registry, pd.DataFrame(files)


def build_fao_screening(registry: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    rows = []
    summaries = []
    for source_key, answers in PRELIMINARY_SCREENING.items():
        score = 0
        for (question_id, question, weight), answer in zip(FAO_SCREENING_QUESTIONS, answers, strict=True):
            awarded = weight if answer == "yes" else 0
            score += awarded
            rows.append({"source_key": source_key, "question_id": question_id, "question": question, "weight": weight, "answer": answer, "awarded_points": awarded, "assessment_type": "preliminary_software_assisted", "requires_expert_confirmation": True})
        summaries.append({"source_key": source_key, "preliminary_score": score, "screening_threshold": 100, "preliminary_pass": score >= 100, "official_full_evaluation_status": "pending_manual_expert_assessment", "rank_eligible": False})
    return pd.DataFrame(rows), pd.DataFrame(summaries)


def build_fao_full_evaluation_queue(root: Path, registry: pd.DataFrame) -> pd.DataFrame:
    """Expand the official 2023 full framework into a source-by-item review queue."""
    workbook = root / "data/raw/reference/standards/FAO_INFOODS_EvaluationFramework_2023.xlsm"
    if not workbook.exists():
        raise FileNotFoundError(f"FAO/INFOODS full evaluation workbook is missing: {workbook}")
    sheet = load_workbook(workbook, read_only=True, data_only=True)["Annex 2.Full evaluation"]
    questions = []
    for row in sheet.iter_rows(min_row=1, values_only=True):
        item_id = str(row[0]).strip() if row[0] is not None else ""
        question = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
        if not re.fullmatch(r"\d+(?:\.\d+){2,}", item_id) or not question:
            continue
        options = [str(value).strip() for value in row[4:10] if value not in (None, "")]
        questions.append({
            "framework_item_id": item_id,
            "target_group": row[1] if len(row) > 1 else "",
            "criterion": question,
            "information": row[3] if len(row) > 3 else "",
            "validation_options": " | ".join(options),
            "priority": row[10] if len(row) > 10 else "",
            "max_score_user": row[11] if len(row) > 11 else "",
            "max_score_compiler": row[12] if len(row) > 12 else "",
        })
    source_rows = registry[registry["local_status"].eq("available") & registry["source_key"].ne("")]
    queue = []
    for source in source_rows.itertuples(index=False):
        for question in questions:
            queue.append({
                "source_key": source.source_key,
                "database_name": source.database_name,
                **question,
                "reviewer_1_answer": "",
                "reviewer_1_evidence": "",
                "reviewer_2_answer": "",
                "reviewer_2_evidence": "",
                "adjudicated_answer": "",
                "adjudication_note": "",
                "review_status": "pending_two_expert_reviewers",
            })
    return pd.DataFrame(queue)


def crosswalk_fao_directory(directory: pd.DataFrame, registry: pd.DataFrame) -> pd.DataFrame:
    review_names: dict[str, list[str]] = {}
    for row in registry.itertuples(index=False):
        key = normalize_text(row.database_name)
        review_names.setdefault(key, []).append(row.registry_id)
    rows = []
    for row in directory.itertuples(index=False):
        stripped = re.sub(r"^\([^)]*\)\s*", "", str(row.title)).strip()
        key = normalize_text(stripped)
        matches = review_names.get(key, [])
        rows.append({
            "fao_directory_id": row.fao_directory_id,
            "fao_title": row.title,
            "normalized_title_without_country_year": key,
            "review_registry_ids": ";".join(matches),
            "crosswalk_status": "exact_title_candidate" if matches else "unmatched_requires_review",
            "automatic_entity_merge": False,
            "review_reason": "Even exact titles remain catalogue-level crosswalk candidates until version and publisher identity are checked.",
        })
    return pd.DataFrame(rows)


def _write_prisma_svg(path: Path, counts: dict[str, int]) -> None:
    boxes = [
        (40, 30, 620, 72, f"Records identified through PubMed, Crossref and OpenAlex: {counts['retrieved_records']:,}"),
        (40, 130, 620, 72, f"Unique candidate records after identifier/title deduplication: {counts['unique_candidates']:,}"),
        (40, 230, 620, 72, f"Curated standards and ontology records included: {counts['curated_included']:,}"),
        (40, 330, 620, 92, "Dual-reviewer title/abstract and full-text screening: PENDING\nNo automated record is treated as included evidence"),
    ]
    parts = ['<svg xmlns="http://www.w3.org/2000/svg" width="700" height="465" viewBox="0 0 700 465">', '<rect width="700" height="465" fill="white"/>']
    for i, (x, y, w, h, label) in enumerate(boxes):
        parts.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="4" fill="#f7f8fa" stroke="#20262e" stroke-width="1.5"/>')
        lines = label.split("\n")
        start = y + h / 2 - (len(lines) - 1) * 10
        for j, line in enumerate(lines):
            parts.append(f'<text x="{x+w/2}" y="{start+j*22}" text-anchor="middle" font-family="Arial" font-size="15" fill="#111820">{line}</text>')
        if i < len(boxes) - 1:
            parts.append(f'<line x1="350" y1="{y+h}" x2="350" y2="{boxes[i+1][1]}" stroke="#20262e" stroke-width="1.5" marker-end="url(#arrow)"/>')
    parts.insert(1, '<defs><marker id="arrow" markerWidth="10" markerHeight="10" refX="5" refY="3" orient="auto"><path d="M0,0 L0,6 L6,3 z" fill="#20262e"/></marker></defs>')
    parts.append('</svg>')
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(parts))


def run(root: Path, out_dir: Path, offline: bool = False) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    logs, records = build_search_corpus(out_dir, offline)
    citation_logs, citation_records = build_citation_tracking(out_dir, offline)
    logs = pd.concat([logs, citation_logs], ignore_index=True, sort=False)
    records = pd.concat([records, citation_records], ignore_index=True, sort=False)
    fao_directory = build_fao_global_directory(out_dir, offline)
    registry, files = build_source_registry(root)
    fao_crosswalk = crosswalk_fao_directory(fao_directory, registry)
    screening_items, screening_summary = build_fao_screening(registry)
    registry = registry.merge(screening_summary, on="source_key", how="left")
    full_evaluation_queue = build_fao_full_evaluation_queue(root, registry)

    curated = pd.DataFrame(CURATED_REFERENCES)
    curated["screening_status"] = "included_by_protocol"
    curated["screening_reason"] = curated["reason"]
    unique_candidates = int(records["dedup_key"].nunique()) if not records.empty else 0
    counts = {
        "retrieved_records": len(records),
        "unique_candidates": unique_candidates,
        "curated_included": len(curated),
        "fao_directory_records": len(fao_directory),
        "human_screened": 0,
        "full_text_assessed": 0,
        "search_as_of": AS_OF_DATE,
    }

    write_csv(logs, out_dir / "search_log.csv")
    write_csv(records, out_dir / "review_records.csv")
    write_csv(citation_logs, out_dir / "citation_tracking_log.csv")
    write_csv(citation_records, out_dir / "citation_tracking_records.csv")
    write_csv(curated, out_dir / "curated_evidence_ledger.csv")
    write_csv(fao_directory, out_dir / "fao_global_directory_snapshot.csv")
    write_csv(fao_crosswalk, out_dir / "fao_directory_to_2025_review_crosswalk.csv")
    write_csv(registry, out_dir / "source_registry.csv")
    write_csv(files, out_dir / "source_file_manifest.csv")
    write_csv(screening_items, out_dir / "fao_infoods_preliminary_screening.csv")
    write_csv(full_evaluation_queue, out_dir / "fao_infoods_full_evaluation_queue.csv.gz")
    write_json(counts, out_dir / "prisma_counts.json")
    _write_prisma_svg(out_dir / "prisma_flow.svg", counts)

    protocol = f"""# Scoping Review Protocol ({DATASET_VERSION})

Search date: {AS_OF_DATE}. Reporting follows PRISMA-ScR and PRISMA-S where applicable.

## Scope

The review covers FCDB inventories, data-quality frameworks, food matching, unit and expression conversion, food/component ontologies, duplicate handling, and composition-prediction benchmarks. It combines a {len(fao_directory):,}-entry snapshot of the complete FAO global directory, the 101-database landscape registry, reproducible PubMed/Crossref/OpenAlex topic searches, and complete OpenAlex backward/forward citation retrieval for the declared domain-method seeds. Directory-to-review crosswalks and citation-index records remain human-review candidates.

## Screening status

The API search is a candidate corpus, not a completed systematic-review inclusion set. Dual independent title/abstract screening, full-text assessment and disagreement adjudication remain mandatory. Rows that have not completed this process are explicitly labelled `awaiting_dual_reviewer_screening` and are not cited as included evidence.

## Standards position

There is no single universal standard that resolves every FCDB identity, quality and interoperability question. This project uses the most authoritative compatible systems for distinct responsibilities: FAO/INFOODS for FCDB evaluation, matching and component expression; EuroFIR QE-SCIREP for value-level quality; FoodOn/LanguaL/FoodEx2 for food description; INFOODS/CDNO/ChEBI/LIPID MAPS for component identity/classification; and FAIR for provenance and release.

## Review gates

1. Complete duplicate removal across DOI, PMID and normalized title.
2. Two reviewers screen every candidate and record an explicit reason for exclusion.
3. Retrieve and assess full text for all retained records.
4. Perform backward and forward citation searches.
5. Freeze the final evidence set before locked validation is opened.

Automated search records: {len(records):,}; unique candidates: {unique_candidates:,}; FAO directory entries: {len(fao_directory):,}; protocol-curated standards: {len(curated):,}.
"""
    (out_dir / "SCOPING_REVIEW_PROTOCOL.md").write_text(protocol)
    print(f"Wrote scoping review and source registry to {out_dir}")
    print(f"Automated records: {len(records):,}; unique candidates: {unique_candidates:,}; FCDB registry rows: {len(registry):,}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path.cwd())
    parser.add_argument("--output-dir", type=Path, default=Path("data/audits/scientific_food_composition_v1/scoping_review"))
    parser.add_argument("--offline", action="store_true", help="Require cached API responses instead of accessing the network.")
    args = parser.parse_args()
    run(args.root.resolve(), (args.root / args.output_dir).resolve(), args.offline)


if __name__ == "__main__":
    main()
